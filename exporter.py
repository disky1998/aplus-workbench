# -*- coding: utf-8 -*-
"""导出商品链接信息表（xlsx / csv）—— 从输出目录里的 product.json 汇总

不依赖内存里的任务状态，所以历史产出也能随时导出。
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
CELL_ALIGN = Alignment(vertical="top", wrap_text=True)

BASE_COLS = [
    ("sku", "SKU"), ("asin", "ASIN"), ("domain", "站点"), ("url", "采集链接"),
    ("title", "标题"), ("highlights", "商品亮点(小标题)"), ("category", "类目路径"),
    ("nodeId", "最小类目节点ID"), ("brand", "品牌"), ("price", "价格"),
    ("rating", "评分"), ("reviews", "评论数"), ("seller", "卖家"),
    ("bulletsSource", "五点来源"),
]


def _load(p: Path) -> dict:
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}


def collect(out_root: Path, items=None) -> list[dict]:
    """扫描 <sku>/product.json，返回按 SKU 排列的行"""
    rows = []
    if items:
        order = [sku for sku, _a, _d in items]
    else:
        order = sorted(x.name for x in out_root.iterdir()
                       if x.is_dir() and (x / "product.json").exists()) \
            if out_root.is_dir() else []

    for sku in order:
        p = out_root / sku / "product.json"
        if not p.exists():
            continue
        d = _load(p)
        if not d:
            continue
        row = {
            "sku": sku,
            "asin": d.get("asin", ""),
            "domain": d.get("domain", ""),
            "url": d.get("url", ""),
            "title": d.get("title", ""),
            "highlights": d.get("highlights", ""),
            "category": d.get("category", ""),
            "nodeId": d.get("nodeId", ""),
            "brand": d.get("brand", ""),
            "price": d.get("price", ""),
            "rating": d.get("rating", ""),
            "reviews": d.get("reviews", ""),
            "seller": d.get("seller", ""),
            "bulletsSource": d.get("bulletsSource", ""),
            "bullets": d.get("bullets") or [],
            "images": d.get("images") or [],
        }
        rows.append(row)
    return rows


def _headers_and_row(row: dict, max_img: int) -> tuple[list, list]:
    heads = [h for _k, h in BASE_COLS]
    vals = [row.get(k, "") for k, _h in BASE_COLS]
    heads += [f"五点{i}" for i in range(1, 6)]
    vals += [(row["bullets"][i] if i < len(row["bullets"]) else "") for i in range(5)]
    heads += [f"主图{i}" for i in range(1, max_img + 1)]
    vals += [(row["images"][i] if i < len(row["images"]) else "") for i in range(max_img)]
    heads.append("标题字数")
    vals.append(len(row.get("title") or ""))
    return heads, vals


def build_xlsx(rows: list[dict], domain: str = "") -> bytes:
    max_img = max([len(r["images"]) for r in rows] or [1])
    max_img = min(max(max_img, 1), 15)
    heads, _ = _headers_and_row(rows[0] if rows else {}, max_img)

    wb = Workbook()
    ws = wb.active
    ws.title = "商品信息"
    ws.append(heads)
    for i in range(1, len(heads) + 1):
        c = ws.cell(row=1, column=i)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26

    for r in rows:
        _, vals = _headers_and_row(r, max_img)
        ws.append(vals)
        for c in ws[ws.max_row]:
            c.alignment = CELL_ALIGN

    for col in ws.columns:
        w = 12
        for c in col[:300]:
            if c.value is None:
                continue
            longest = max((len(seg) for seg in str(c.value).split("\n")), default=0)
            w = max(w, longest * 1.05 + 2)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(58, w)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(rows: list[dict]) -> bytes:
    """保留：CSV 的列与 Excel 完全一致，已由 build_xlsx 取代（界面只留一个入口）"""
    max_img = min(max([len(r["images"]) for r in rows] or [1]), 15)
    heads, _ = _headers_and_row(rows[0] if rows else {}, max_img)
    sio = io.StringIO()
    w = csv.writer(sio, quoting=csv.QUOTE_ALL, lineterminator="\n")
    w.writerow(heads)
    for r in rows:
        _, vals = _headers_and_row(r, max_img)
        w.writerow(vals)
    return ("\ufeff" + sio.getvalue()).encode("utf-8")


# 前 8 列顺序必须与「批量文案重构」工具的 Excel 导入要求严格一致：
#   SKU / 原标题 / 原亮点 / 五点1..五点5
# 后面的列是溯源信息，导入时会被工具忽略。
TOOL_HEAD = ["SKU", "原标题", "原亮点", "五点1", "五点2", "五点3", "五点4", "五点5"]
TOOL_TRACE = ["ASIN", "站点", "采集链接", "类目路径", "最小类目节点ID", "品牌",
              "价格", "评分", "评论数", "五点来源"]


def build_tool_xlsx(rows: list[dict]) -> bytes:
    """可直接上传到「批量文案重构」的导入表（xlsx）"""
    max_img = min(max([len(r["images"]) for r in rows] or [1]), 15)
    heads = list(TOOL_HEAD) + list(TOOL_TRACE) + [f"原主图{i}" for i in range(1, max_img + 1)]

    wb = Workbook()
    ws = wb.active
    ws.title = "批量改写导入"
    ws.append(heads)
    for i in range(1, 9):                    # 前 8 列标成重点色，方便一眼核对
        c = ws.cell(row=1, column=i)
        c.fill = PatternFill("solid", fgColor="C55A11")
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for i in range(9, len(heads) + 1):
        c = ws.cell(row=1, column=i)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26

    for r in rows:
        bl = r.get("bullets") or []
        vals = [r.get("sku", ""), r.get("title", ""), r.get("highlights", "")]
        vals += [bl[i] if i < len(bl) else "" for i in range(5)]
        vals += [r.get(k, "") for k in ("asin", "domain", "url", "category",
                                        "nodeId", "brand", "price", "rating",
                                        "reviews", "bulletsSource")]
        imgs = r.get("images") or []
        vals += [imgs[i] if i < len(imgs) else "" for i in range(max_img)]
        ws.append(vals)
        for c in ws[ws.max_row]:
            c.alignment = CELL_ALIGN

    for col in ws.columns:
        w = 12
        for c in col[:300]:
            if c.value is None:
                continue
            longest = max((len(seg) for seg in str(c.value).split("\n")), default=0)
            w = max(w, longest * 1.05 + 2)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(58, w)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
