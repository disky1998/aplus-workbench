# -*- coding: utf-8 -*-
"""验证导出接口 + 真实抓取 all 模式（商品信息 + A+ 一起）"""
import json
import time
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
BASE = "http://127.0.0.1:8799"
out = {}


def call(path, method="GET", body=None, timeout=120, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(BASE + path, data=data, method=method,
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        b = r.read()
        if raw:
            return {"status": r.status, "bytes": len(b),
                    "cd": r.headers.get("Content-Disposition"),
                    "ctype": r.headers.get("Content-Type"), "_b": b}
        return json.loads(b.decode("utf-8"))


for _ in range(40):
    try:
        call("/api/version", timeout=5); break
    except Exception:
        time.sleep(0.5)

# 1) 已有产出的导出
for kind, path in (("xlsx", "/api/export/product.xlsx?dir=20260910-150107"),
                   ("csv", "/api/export/product.csv?dir=20260910-150107")):
    try:
        r = call(path, raw=True)
        info = {k: r[k] for k in ("status", "bytes", "cd", "ctype")}
        if kind == "xlsx":
            p = HERE / "_chk.xlsx"
            p.write_bytes(r["_b"])
            from openpyxl import load_workbook
            ws = load_workbook(p).active
            info["headers"] = [c.value for c in ws[1]][:8]
            info["row1"] = [str(c.value)[:40] for c in ws[2]][:8]
            p.unlink()
        out[f"export_{kind}"] = info
    except Exception as e:
        out[f"export_{kind}"] = {"error": f"{type(e).__name__}: {e}"}

# 2) 真实抓取 all 模式（商品信息 + A+）
try:
    r = call("/api/run", "POST", {
        "text": "ALL-TEST https://www.amazon.ae/dp/B0BTPX24RV",
        "domain": "ae", "want": "all", "both": True, "aplusMode": "auto",
        "downloadImages": True, "includeBrandStory": False, "screenshot": False,
        "hires": 1464, "delay": 0, "retry": 0, "headed": False,
        "proxy": "", "userDataDir": "", "persist": True,
        "buildCombined": True, "packageImages": True}, timeout=60)
    tid = r["task_id"]
    since, logs, final = 0, [], None
    deadline = time.time() + 300
    while time.time() < deadline:
        st = call(f"/api/task/{tid}?since={since}", timeout=60)
        logs += st.get("logs", [])
        since = st.get("next", since)
        if st["status"] in ("done", "error", "stopped"):
            final = st; break
        time.sleep(2)
    res = (final or {}).get("results") or [{}]
    out["live_all"] = {
        "task": tid, "status": (final or {}).get("status"),
        "combined": (final or {}).get("combined"),
        "zipfile": (final or {}).get("zipfile"),
        "sku": res[0].get("sku"), "ok": res[0].get("ok"),
        "kind": res[0].get("kind"),
        "product": res[0].get("product"),
        "viewports": res[0].get("viewports"),
        "logs": logs[-20:],
    }
    if final and final.get("combined"):
        d = final["combined"].split("/")[0]
        out["live_all"]["export_after"] = call(
            f"/api/export/product.xlsx?dir={d}", raw=True)["bytes"]
except Exception as e:
    out["live_all"] = {"error": f"{type(e).__name__}: {e}"}

out["history"] = [{"id": it["id"], "has_product": it["has_product"],
                   "product_count": it["product_count"],
                   "has_combined": it["has_combined"], "has_zip": it["has_zip"]}
                  for it in call("/api/history")["items"][:3]]

HERE.joinpath("_e2e3.json").write_text(json.dumps(out, ensure_ascii=False, indent=2),
                                       encoding="utf-8")
print("done")
