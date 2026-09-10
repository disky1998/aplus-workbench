# -*- coding: utf-8 -*-
"""离线验证：商品信息解析（模拟页）+ 导出接口 + 页面元素"""
import json
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
out = {}

FAKE_HTML = """<!DOCTYPE html><html><body>
<div id="wayfinding-breadcrumbs_feature_div"><ul>
  <li><a class="a-link-normal" href="/b?node=1">Home</a></li>
  <li><a class="a-link-normal" href="/b?node=2">Kitchen &amp; Dining</a></li>
  <li><a class="a-link-normal" href="/b?node=15172727031">Gift Bags</a></li>
</ul></div>
<span id="productTitle">   Servewell  Harmony  Gift Bag ,  Multicolor   </span>
<div id="itemHighlights_feature_div"><ul>
  <li>Material: Paper</li><li>Colour: Assorted</li><li>Occasion: Birthday</li></ul></div>
<div id="feature-bullets"><ul>
  <li><span class="a-list-item">Made from sturdy paper with a glossy finish for a premium look</span></li>
  <li><span class="a-list-item">Assorted designs ship randomly, each bag measures 26 x 32 cm</span></li>
  <li><span class="a-list-item">Click here to report an issue with this product</span></li>
</ul></div>
<img id="landingImage" src="https://m.media-amazon.com/images/I/81eVtzRgQmL._AC_SX679_.jpg"
     data-a-dynamic-image='{"https://m.media-amazon.com/images/I/81eVtzRgQmL._AC_SX679_.jpg":[679,679]}'>
<div id="altImages"><ul>
  <li><img src="https://m.media-amazon.com/images/I/51NQXKYhh3L._AC_UL320_.jpg"></li>
  <li><img src="https://m.media-amazon.com/images/I/transparent-pixel._AC_UL320_.jpg"></li>
</ul></div>
<a id="bylineInfo">Brand: Servewell</a>
<div id="corePriceDisplay_desktop_feature_div">
  <span class="a-price"><span class="a-offscreen">AED5.00</span></span></div>
<div id="acrPopover" title="3.8 out of 5 stars"></div>
<span id="acrCustomerReviewText">75 ratings</span>
<pre id="prod')
</pre>
</body></html>"""

FAKE_HTML = FAKE_HTML.replace("<pre id=\"prod')\n</pre>\n", "")


def main():
    import product as prod
    from playwright.sync_api import sync_playwright

    tmp = HERE / "_fake_page.html"
    tmp.write_text(FAKE_HTML, encoding="utf-8")

    with sync_playwright() as p:
        br = p.chromium.launch(headless=True)
        page = br.new_context(locale="en-US").new_page()
        page.goto(tmp.as_uri(), wait_until="domcontentloaded")
        data = prod.extract_product(page)
        br.close()
    tmp.unlink(missing_ok=True)

    out["parsed"] = {k: data.get(k) for k in
                     ("title", "highlights", "category", "nodeId", "bullets",
                      "images", "brand", "price", "rating", "reviews", "bulletsSource")}
    out["missing"] = prod.missing_fields(data)
    out["checks"] = {
        "标题归一化": data.get("title") == "Servewell Harmony Gift Bag , Multicolor",
        "亮点合并": (data.get("highlights") or "").startswith("Material: Paper, Colour: Assorted"),
        "类目路径": data.get("category") == "Home > Kitchen & Dining > Gift Bags",
        "NodeID": data.get("nodeId") == "15172727031",
        "五点噪音已过滤": len(data.get("bullets") or []) == 2,
        "主图还原master": data.get("images") == [
            "https://m.media-amazon.com/images/I/81eVtzRgQmL.jpg",
            "https://m.media-amazon.com/images/I/51NQXKYhh3L.jpg"],
        "品牌去前缀": data.get("brand") == "Servewell",
        "价格": data.get("price") == "AED5.00",
        "评分": data.get("rating") == "3.8",
        "评论数": data.get("reviews") == "75",
    }

    # ---- 导出（造一个假的 product.json 目录）----
    import webapp
    d = webapp.WEB_OUT / "_test-export"
    sku_dir = d / "SKU-A"
    sku_dir.mkdir(parents=True, exist_ok=True)
    rec = dict(data)
    rec.update({"asin": "B0TEST00001", "domain": "ae",
                "url": "https://www.amazon.ae/dp/B0TEST00001"})
    (sku_dir / "product.json").write_text(json.dumps(rec, ensure_ascii=False, indent=2),
                                          encoding="utf-8")

    import exporter as exp
    rows = exp.collect(d)
    xlsx = exp.build_xlsx(rows)
    csvb = exp.build_csv(rows)
    (HERE / "_t_product.xlsx").write_bytes(xlsx)

    from openpyxl import load_workbook
    wb = load_workbook(HERE / "_t_product.xlsx")
    ws = wb.active
    heads = [c.value for c in ws[1]]
    row1 = [c.value for c in ws[2]]
    out["export"] = {
        "rows": len(rows),
        "xlsx_bytes": len(xlsx),
        "csv_bytes": len(csvb),
        "headers": heads,
        "row1_head": row1[:8],
        "五点列": [row1[heads.index(f"五点{i}")] for i in (1, 2, 3)],
        "主图列": [row1[heads.index(f"主图{i}")] for i in (1, 2)],
    }
    for p in (HERE / "_t_product.xlsx",):
        p.unlink(missing_ok=True)
    import shutil
    shutil.rmtree(d, ignore_errors=True)

    (HERE / "_prodt.json").write_text(json.dumps(out, ensure_ascii=False, indent=2),
                                      encoding="utf-8")
    print("done")


if __name__ == "__main__":
    main()
